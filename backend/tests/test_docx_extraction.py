import pytest
from app.services.ingestion_service import IngestionService
import docx
import os
from unittest.mock import MagicMock

def test_extract_docx_with_headings_and_tables(tmp_path):
    # Create a mock docx file
    doc_path = tmp_path / "test.docx"
    doc = docx.Document()
    
    # Add Heading 1
    doc.add_heading("Employee Benefits", level=1)
    # Add paragraph
    doc.add_paragraph("The company provides health insurance.")
    # Add Heading 2
    doc.add_heading("Health Insurance", level=2)
    # Add table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Plan"
    table.cell(0, 1).text = "Coverage"
    table.cell(1, 0).text = "Gold"
    table.cell(1, 1).text = "Family"
    
    doc.save(doc_path)
    
    # Create service
    service = IngestionService(db=MagicMock(), embedder=MagicMock(), lexical_store=None)
    
    # Extract
    text = service._extract_docx(str(doc_path))
    
    assert "# Employee Benefits" in text
    assert "The company provides health insurance." in text
    assert "## Health Insurance" in text
    assert "[TABLE]" in text
    assert "Plan | Coverage" in text
    assert "Gold | Family" in text
